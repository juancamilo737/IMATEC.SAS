#!/usr/bin/env python3
"""Extrae los ítems de las cotizaciones históricas de IMATEC.

Cada archivo de «COTIZACIONES IMATEC» usa la misma plantilla: una fila de
cabecera con ITEM / UNIDAD / VR. UNIDAD / CANTIDAD / VR. TOTAL y debajo las
líneas, hasta la fila que empieza con «SON :».

Devuelve un consolidado por producto: cuántas veces se cotizó, en qué
unidades y a qué precios. Es el histórico real de venta de la empresa.
"""
import glob
import json
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime

from openpyxl import load_workbook

CARPETA = "Archivos/COTIZACIONES IMATEC/*.xlsx"


def norm(t) -> str:
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    t = t.upper().replace("*", '"')
    t = re.sub(r"[^A-Z0-9/\"\.\- ]", " ", t)
    return " ".join(t.split())


def fila_cabecera(ws):
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, 9)]
        if "ITEM" in vals and any("VR. UNIDAD" in v or "VR UNIDAD" in v for v in vals):
            return r
    return None


def extraer(ruta):
    """Ítems de una cotización. Devuelve (fecha, cliente, [items])."""
    try:
        wb = load_workbook(ruta, data_only=True, read_only=True)
    except Exception:
        return None, None, []
    ws = wb.worksheets[0]
    cab = fila_cabecera(ws)
    if not cab:
        wb.close()
        return None, None, []

    cliente = fecha = None
    for r in range(1, cab):
        etiqueta = norm(ws.cell(r, 1).value)
        if "NOMBRE CLIENTE" in etiqueta:
            cliente = ws.cell(r, 3).value
        for c in range(1, 9):
            if norm(ws.cell(r, c).value) == "FECHA":
                v = ws.cell(r, c + 2).value or ws.cell(r, c + 1).value
                if isinstance(v, datetime):
                    fecha = v.date().isoformat()

    items = []
    for r in range(cab + 1, ws.max_row + 1):
        primera = norm(ws.cell(r, 1).value)
        if primera.startswith("SON"):
            break
        desc = ws.cell(r, 2).value
        precio = ws.cell(r, 6).value
        cant = ws.cell(r, 7).value
        if not desc or not isinstance(precio, (int, float)) or precio <= 0:
            continue
        items.append({
            "descripcion": " ".join(str(desc).split()),
            "unidad": " ".join(str(ws.cell(r, 5).value or "").split()),
            "precio": float(precio),
            "cantidad": float(cant) if isinstance(cant, (int, float)) else 1.0,
        })
    wb.close()
    return fecha, (str(cliente).strip() if cliente else None), items


def main():
    archivos = sorted(glob.glob(CARPETA))
    print(f"Leyendo {len(archivos)} cotizaciones…", flush=True)

    prod = defaultdict(lambda: {"descripcion": "", "veces": 0, "unidades": defaultdict(int),
                                "precios": [], "ultima": "", "cantidad_total": 0.0})
    leidas = vacias = 0
    for i, ruta in enumerate(archivos, 1):
        if i % 200 == 0:
            print(f"  {i}/{len(archivos)}…", flush=True)
        fecha, _cliente, items = extraer(ruta)
        if not items:
            vacias += 1
            continue
        leidas += 1
        for it in items:
            clave = norm(it["descripcion"])
            if len(clave) < 4:
                continue
            p = prod[clave]
            p["descripcion"] = p["descripcion"] or it["descripcion"]
            p["veces"] += 1
            p["cantidad_total"] += it["cantidad"]
            if it["unidad"]:
                p["unidades"][it["unidad"]] += 1
            p["precios"].append((fecha or "", it["precio"]))
            if fecha and fecha > p["ultima"]:
                p["ultima"] = fecha

    salida = []
    for clave, p in prod.items():
        precios = sorted(p["precios"])
        valores = [v for _, v in precios]
        salida.append({
            "clave": clave,
            "descripcion": p["descripcion"],
            "veces": p["veces"],
            "cantidad_total": round(p["cantidad_total"], 2),
            "unidad": max(p["unidades"], key=p["unidades"].get) if p["unidades"] else "",
            "precio_ultimo": precios[-1][1] if precios else 0,
            "precio_min": min(valores) if valores else 0,
            "precio_max": max(valores) if valores else 0,
            "ultima_fecha": p["ultima"],
        })
    salida.sort(key=lambda x: -x["veces"])

    print(f"\nCotizaciones con ítems: {leidas} | sin ítems o ilegibles: {vacias}")
    print(f"Productos distintos:    {len(salida)}")
    print(f"Líneas totales:         {sum(p['veces'] for p in salida)}")
    with open("Archivos/items_consolidados.json", "w", encoding="utf-8") as f:
        json.dump(salida, f, ensure_ascii=False, indent=1)
    print("Guardado en Archivos/items_consolidados.json")

    print("\nLos 15 productos más cotizados:")
    for p in salida[:15]:
        print(f"  {p['veces']:>4}x  {p['descripcion'][:44]:46s} {p['unidad'][:12]:14s} "
              f"${p['precio_ultimo']:>11,.0f}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
